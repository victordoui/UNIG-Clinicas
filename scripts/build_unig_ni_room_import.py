"""Gera SQL idempotente de importação do inventário de espaços UNIG-NI."""
from __future__ import annotations

import re
import sys
import unicodedata
from pathlib import Path
from openpyxl import load_workbook


def slug(value: object) -> str:
    normalized = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^A-Za-z0-9]+", "-", normalized.upper()).strip("-")
    return text or "ESPACO"


def yes(value: object) -> bool:
    return str(value or "").strip() == "✅"


def room_type(name: str) -> str:
    normalized = name.lower()
    if "audit" in normalized:
        return "auditorio"
    if "lab" in normalized or "informática" in normalized or "informatica" in normalized:
        return "laboratorio"
    if "biblioteca" in normalized:
        return "biblioteca"
    if "metodologia" in normalized:
        return "sala_metodologia"
    if "desenho" in normalized:
        return "sala_especial"
    return "sala_aula"


def quality(value: object) -> str:
    text = str(value or "").lower()
    if "premium" in text and "semi" not in text:
        return "premium"
    if "semi" in text:
        return "semi_premium"
    return "padrao"


def sql_text(value: object | None) -> str:
    if value is None or str(value).strip() == "":
        return "null"
    return "'" + str(value).replace("'", "''").strip() + "'"


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Uso: build_unig_ni_room_import.py <entrada.xlsx> <saida.sql>")
    source, output = map(Path, sys.argv[1:])
    workbook = load_workbook(source, data_only=True, read_only=True)
    sheet = workbook.active
    rows = []
    for row in sheet.iter_rows(min_row=3, values_only=True):
        block, room, capacity, chairs, projector, audio, tv, air, board, interactive, tier, furniture, status, usage = row[:14]
        if not room or not isinstance(capacity, (int, float)):
            continue
        room_name = str(room).strip()
        block_name = str(block).strip()
        code = f"{slug(block_name)}-{slug(room_name)}"
        status_value = "disponivel" if "dispon" in str(status or "").lower() or status is None else "inativa"
        notes = f"Importado da planilha de capacidade UNIG-NI. Tipo original: {str(tier or 'Não informado').strip()}."
        rows.append((code, room_name, block_name, int(capacity), int(chairs) if isinstance(chairs, (int, float)) else None, room_type(room_name), yes(projector), yes(air), yes(audio), yes(tv), yes(board), yes(interactive), str(furniture).strip() if furniture else None, quality(tier), str(usage).strip() if usage else None, status_value, notes))

    values = []
    for item in rows:
        code, name, block, capacity, seats, kind, projector, air, audio, tv, board, interactive, furniture, tier, usage, status, notes = item
        values.append("(" + ", ".join([sql_text(code), sql_text(name), sql_text(block), "null", str(capacity), sql_text(kind), str(projector).lower(), str(air).lower(), "false", sql_text(status), sql_text(notes), str(seats) if seats is not None else "null", str(audio).lower(), str(tv).lower(), str(board).lower(), str(interactive).lower(), sql_text(furniture), sql_text(tier), sql_text(usage)]) + ")")

    output.write_text("""-- Importação idempotente: inventário de espaços do Campus Nova Iguaçu.\n-- Origem: Controle de Capacidade de Salas e Mapa UNIG Nova Iguaçu.xlsx\n-- Pré-requisito: 20260813025730_room_inventory_metadata.sql aplicado.\n\nwith target_unit as (\n  select id from public.units where code = 'UNIG-NI' and is_active = true limit 1\n), source_rooms (code, name, block, floor, capacity, room_type, has_projector, has_air_conditioning, has_computer, status, notes, seats_count, has_audio_system, has_tv, has_whiteboard, has_interactive_screen, furniture_type, quality_tier, usage_restriction) as (\n  values\n    """ + ",\n    ".join(values) + """\n)\ninsert into public.rooms (unit_id, code, name, block, floor, capacity, room_type, has_projector, has_air_conditioning, has_computer, status, notes, seats_count, has_audio_system, has_tv, has_whiteboard, has_interactive_screen, furniture_type, quality_tier, usage_restriction)\nselect target_unit.id, source_rooms.*\nfrom target_unit cross join source_rooms\nwhere not exists (\n  select 1 from public.rooms existing\n  where existing.unit_id = target_unit.id and existing.code = source_rooms.code\n);\n\nselect count(*) as imported_or_existing_rooms\nfrom public.rooms\nwhere unit_id = (select id from public.units where code = 'UNIG-NI' limit 1);\n""", encoding="utf-8")


if __name__ == "__main__":
    main()
